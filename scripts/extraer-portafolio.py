#!/usr/bin/env python3
"""
Portafolio del cliente (.xlsx, dos pestañas) -> JSON crudo para la carga.

El .mjs de carga vive en el lenguaje del proyecto (Node), pero Node no lee
.xlsx sin dependencias. Este extractor hace SOLO eso: vuelca las filas tal
cual a `scripts/datos/portafolio.json`. Toda la interpretación (tipología ->
materia/brillo/textura, normalizar rendimiento, etc.) la hace el .mjs.

    python3 scripts/extraer-portafolio.py            # ruta por defecto
    XLSX=/otra/ruta.xlsx python3 scripts/extraer-portafolio.py

Requiere `openpyxl`. El dato marcado en AMARILLO en el Excel (dureza,
rectificado y absorción de la Serie Venezuela) es el que el cliente completó
a mano porque no estaba en el catálogo PDF; se marca en `aMano` para poder
armar la lista de verificación.
"""
import json, os, sys
try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instálalo en un venv:\n"
             "  python3 -m venv .venv && .venv/bin/pip install openpyxl\n"
             "  .venv/bin/python scripts/extraer-portafolio.py")

XLSX = os.environ.get(
    "XLSX",
    "/root/fotos-cliente/Portafolio_Ceramica_Carabobo_SerieRegular_SerieVenezuela (REVISADO).xlsx",
)
SALIDA = os.path.join(os.path.dirname(__file__), "datos", "portafolio.json")
AMARILLO = "FFFFFF00"

# Columnas (1-based) de la fila de encabezado (fila 9) hacia abajo.
COL = dict(item=1, nombre=3, formato=4, tipologia=5, pei=6, dureza=7,
           rectificado=8, absorcion=9, rendimiento=10)
PESTANAS = [("Serie Regular", "Regular"), ("Serie Venezuela", "Venezuela")]


def texto(v):
    return None if v is None else str(v).strip()


def es_amarillo(celda):
    f = celda.fill
    return f.patternType == "solid" and str(getattr(f.fgColor, "rgb", "")) == AMARILLO


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    filas = []
    for hoja, serie in PESTANAS:
        ws = wb[hoja]
        for r in range(11, ws.max_row + 1):
            item = ws.cell(r, COL["item"]).value
            nombre = ws.cell(r, COL["nombre"]).value
            if not isinstance(item, (int, float)) or not nombre:
                continue  # filas de título, leyenda o vacías
            filas.append({
                "item": int(item),
                "serie": serie,
                "nombre": texto(nombre),
                "formato": texto(ws.cell(r, COL["formato"]).value),
                "tipologia": texto(ws.cell(r, COL["tipologia"]).value),
                "pei": texto(ws.cell(r, COL["pei"]).value),
                "dureza": ws.cell(r, COL["dureza"]).value,
                "rectificado": texto(ws.cell(r, COL["rectificado"]).value),
                "absorcion": texto(ws.cell(r, COL["absorcion"]).value),
                "rendimiento": texto(ws.cell(r, COL["rendimiento"]).value),
                "aMano": {
                    "dureza": es_amarillo(ws.cell(r, COL["dureza"])),
                    "rectificado": es_amarillo(ws.cell(r, COL["rectificado"])),
                    "absorcion": es_amarillo(ws.cell(r, COL["absorcion"])),
                },
            })
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    with open(SALIDA, "w", encoding="utf-8") as fh:
        json.dump(filas, fh, ensure_ascii=False, indent=1)
    reg = sum(1 for f in filas if f["serie"] == "Regular")
    ven = len(filas) - reg
    amano = sum(1 for f in filas if any(f["aMano"].values()))
    print(f"{SALIDA}: {len(filas)} productos (Regular {reg}, Venezuela {ven}) · {amano} con dato a mano")


if __name__ == "__main__":
    main()
